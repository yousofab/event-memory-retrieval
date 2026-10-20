"""Build the Phase 16 Persian RTL research report from docs/phase16.md."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_phase15_report as prior


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/phase16.md"
DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۱۶-مرزهای-همتعداد.docx"


def build():
    prior.prior.SOURCE = SOURCE
    prior.prior.DESTINATION = DESTINATION
    prior.build()
    doc = Document(DESTINATION)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    prior.prior.prior.base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    prior.prior.prior.base.add_run(footer, "گزارش پژوهشی مرحله‌ی شانزدهم  |  صفحه ", size=10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "گزارش پژوهشی مرحله‌ی شانزدهم مرزهای هم‌تعداد در حافظه‌ی گفت‌وگو"
    doc.core_properties.subject = "اثر جای مرز با کنترل نقش و سقف واژه در Memora"
    doc.save(DESTINATION)
    return DESTINATION


if __name__ == "__main__":
    print(build())
