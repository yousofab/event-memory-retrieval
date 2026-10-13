"""Build the editable Persian RTL Phase 9 report, with the supplied logo.

The logo was provided by the researcher; B Nazanin remains a local font and is
not redistributed. The builder depends on python-docx and build_phase5_report.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

import build_phase5_report as base


ROOT = Path(__file__).resolve().parents[1]
base.SOURCE = ROOT / "docs/phase9.md"
base.DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۹-کنترل-مرزها.docx"
original_setup = base.setup
original_table = base.add_table


def setup(document):
    original_setup(document)
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    base.add_run(footer, "گزارش پژوهشی مرحله‌ی نهم  |  صفحه ", size=10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_table(document, block):
    original_table(document, block)
    table = document.tables[-1]
    for i, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if i == 0:
            properties.append(OxmlElement("w:tblHeader"))


def add_hyperlink(paragraph, url):
    relation = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "DejaVu Sans")
    fonts.set(qn("w:hAnsi"), "DejaVu Sans")
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "123F5A")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    run.append(properties)
    label = OxmlElement("w:t")
    label.text = "source"
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def build():
    base.setup = setup
    base.add_table = add_table
    base.build()
    document = Document(base.DESTINATION)
    for paragraph in document.paragraphs:
        if paragraph.text == "ناهمگنی گفت‌وگوها":
            paragraph.paragraph_format.page_break_before = True
            break
    logo = ROOT / "assets/fardiat-logo.jpg"
    if logo.exists():
        for i, paragraph in enumerate(list(document.paragraphs)):
            if paragraph.text == "سپاسگزاری و خاستگاه آموزشی":
                target = document.paragraphs[i + 1]
                figure = document.add_paragraph()
                base.persian_paragraph(figure, WD_ALIGN_PARAGRAPH.CENTER)
                figure.paragraph_format.space_before = Pt(4)
                figure.paragraph_format.space_after = Pt(6)
                figure.add_run().add_picture(str(logo), width=Cm(3.0))
                target._p.addprevious(figure._p)
                break
    references, in_references = [], False
    for line in base.SOURCE.read_text(encoding="utf-8").splitlines():
        if line.strip() == "## منابع":
            in_references = True
        elif in_references and line.strip().startswith("["):
            if " https://" in line:
                citation, url = line.rsplit(" https://", 1)
                references.append((citation, "https://" + url))
            else:
                references.append((line, None))
    in_references = False
    for paragraph in list(document.paragraphs):
        if paragraph.text == "منابع":
            in_references = True
            continue
        if in_references:
            paragraph._element.getparent().remove(paragraph._element)
    for citation, url in references:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.06
        run = paragraph.add_run(citation + "  ")
        run.font.name = "DejaVu Sans"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if url:
            add_hyperlink(paragraph, url)
    document.core_properties.title = "گزارش پژوهشی مرحله‌ی نهم کنترل تصادفی مرزها"
    document.core_properties.subject = "مرزبندی حافظه‌ی گفت‌وگو با کنترل شمار قطعه"
    document.save(base.DESTINATION)
    return base.DESTINATION


if __name__ == "__main__":
    print(build())
