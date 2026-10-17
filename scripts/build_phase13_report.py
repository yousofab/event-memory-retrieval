"""Build the Phase 13 Persian RTL Word report from registered Markdown."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_phase9_report as prior


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/phase13.md"
DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۱۳-بازیابی-شاهد-از-دو-جلسه.docx"


def build():
    prior.base.SOURCE = SOURCE
    prior.base.DESTINATION = DESTINATION
    prior.build()
    doc = Document(DESTINATION)
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 1.06
        if paragraph.text.startswith(("[۱]", "[۲]", "[۳]", "[۴]", "[۵]", "[۶]")):
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
    footer = section.footer.paragraphs[0]
    footer.clear()
    prior.base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    prior.base.add_run(footer, "گزارش پژوهشی مرحله‌ی سیزدهم  |  صفحه ", size=10)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    doc.core_properties.title = "گزارش پژوهشی مرحله‌ی سیزدهم بازیابی شاهد از دو جلسه"
    doc.core_properties.subject = "ممیزی تنوع جلسه در بازیابی شواهد پیش و پس از اصلاح اطلاعات"
    doc.save(DESTINATION)
    return DESTINATION


if __name__ == "__main__":
    print(build())
