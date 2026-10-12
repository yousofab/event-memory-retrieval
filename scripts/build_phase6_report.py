"""Build the Persian RTL Phase 6 DOCX from its Markdown source.

Requires python-docx and a locally installed B Nazanin font for the intended
render. Neither a font file nor raw LoCoMo text is committed to this project.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

import build_phase5_report as base


ROOT = Path(__file__).resolve().parents[1]
base.SOURCE = ROOT / "docs/phase6.md"
base.DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۶-مرزبندی-معنایی.docx"
original_setup = base.setup
original_table = base.add_table


def setup(document):
    original_setup(document)
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    base.add_run(footer, "گزارش پژوهشی مرحله‌ی ششم  |  صفحه ", size=10)
    page_number = OxmlElement("w:fldSimple")
    page_number.set(qn("w:instr"), "PAGE")
    footer._p.append(page_number)


def add_table(document, block):
    original_table(document, block)
    table = document.tables[-1]
    for i, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if i == 0:
            properties.append(OxmlElement("w:tblHeader"))


def add_hyperlink(paragraph, url):
    relation = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "DejaVu Sans")
    fonts.set(qn("w:hAnsi"), "DejaVu Sans")
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "123F5A")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    run.append(properties)
    label = OxmlElement("w:t")
    label.text = "source"
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def build():
    base.setup = setup
    base.add_table = add_table
    base.build()
    doc = Document(base.DESTINATION)
    reference_lines = []
    in_references = False
    for line in base.SOURCE.read_text(encoding="utf-8").splitlines():
        if line.strip() == "## منابع":
            in_references = True
        elif in_references and line.strip().startswith("["):
            citation, url = line.rsplit(" https://", 1)
            reference_lines.append((citation, "https://" + url))
    in_references = False
    for paragraph in list(doc.paragraphs):
        if paragraph.text == "منابع":
            in_references = True
            paragraph.paragraph_format.page_break_before = True
            continue
        if in_references:
            paragraph._element.getparent().remove(paragraph._element)
    for citation, url in reference_lines:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.06
        run = paragraph.add_run(citation + "  ")
        run.font.name = "DejaVu Sans"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)
        add_hyperlink(paragraph, url)
    doc.core_properties.title = "گزارش پژوهشی مرحله‌ی ششم مرزبندی معنایی اپیزودها و بازیابی شواهد"
    doc.core_properties.subject = "مقایسه‌ی اکتشافی LSA و مرز واژگانی در LoCoMo"
    doc.save(base.DESTINATION)
    return base.DESTINATION


if __name__ == "__main__":
    print(build())
