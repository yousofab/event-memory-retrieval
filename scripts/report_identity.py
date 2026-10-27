"""Shared academic identity block for current and future Persian reports."""
import json
from pathlib import Path
from docx.shared import Inches,Pt
from package_phase17 import rtl

ROOT=Path(__file__).resolve().parents[1]


def university_mark(doc, identity_path='configs/report_identity.json'):
    # The default preserves historical Phase19 reproduction. New thesis reports
    # explicitly select the corrected current identity snapshot.
    settings=json.loads((ROOT/identity_path).read_text())
    logo=ROOT/settings['university_logo']
    if not logo.is_file():
        raise FileNotFoundError('The university logo is required for this academic report')
    p=doc.add_paragraph();rtl(p)
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=True
    run=p.add_run();run.add_picture(str(logo),width=Inches(.62))
    for drawing in run._r.xpath('.//wp:docPr'):
        drawing.set('descr',settings['university_fa'])
    return settings
