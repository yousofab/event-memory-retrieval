"""Build the supervisor-revised thesis with IEEE citations and editable mathematics."""
import argparse
import json
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import package_thesis_integrated as old
from report_identity import university_mark
from build_thesis_synthesis import CHAPTERS
from thesis_ieee import References

ROOT=Path(__file__).resolve().parents[1]
FOLDER=ROOT/'thesis'
FA='B Nazanin'
EN='Nimbus Roman'
REFS=None
BOOKMARK=1000
FIGURES=[
 ('01_pipeline','زنجیره‌ی کنترل‌شده‌ی بازیابی و محل ورود برچسب مرجع','این طرح ساختار محاسباتی را نشان می‌دهد؛ برچسب مرجع فقط وارد ممیزی می‌شود [S17، S18F].'),
 ('02_controlled_effects','اندازه‌ی اثر راهنمای مرز در فرصت‌های ثابت','خطوط افقی بازه‌های توصیفی خوشه‌ای ۹۵٪ هستند. خط چین دو واحد درصد آستانه‌ی عملی مرحله‌ی ۱۷ است. سنجه‌ی دو پنل متفاوت است [C14].'),
 ('03_qmsum_gap','شکاف هم‌راستایی مرز و پوشش بازیابی در QMSum','پنل چپ سنجه‌ی آغاز موضوع و پنل راست پوشش نوبت مرتبط را نشان می‌دهد. نامزدهای پنل راست چهار برابر خروجی بودجه دارند؛ فاصله‌ی آن‌ها اثر خالص یک رتبه‌بند نیست [C01 تا C05].'),
 ('04_failure_decomposition','افراز تشخیصی پرسش‌ها در بودجه‌ی ۱۰۲۴ واژه','طول نوار سهم پرسش‌ها و عدد داخل آن شمار پرسش است. صفرِ جلسه‌ی غایب QMSum بخش رنگی ندارد و خطای انتخاب ممکن آن ۶ مورد است. سقف LoCoMo برابر ۴۶۸ و سقف QMSum برابر ۲۹ مورد است [C07 تا C09].'),
 ('05_access_assembly_locomo','دسترسی و امکان و پوشش در LoCoMo','تعدادها برای ۵۹۹ پرسش، نویسنده‌ی ثابت و بودجه‌ی ۱۰۲۴ واژه‌اند. افزایش دسترسی از ۴۹۳ به ۵۰۷ با افت امکان از ۴۶۸ به ۴۶۱ و پوشش از ۴۴۳ به ۴۰۱ همراه است [C11].'),
 ('06_access_assembly_external','دسترسی و پوشش پیام مرجع در منبع بیرونی','پنل چپ تعداد پرسش از ۶۰۹ و پنل راست میانگین مساوی چهار پروژه برای کسر پیام مرجع کامل است. این دو پنل مخرج و نوع تجمیع متفاوت دارند [C12].')]
TABLES=[]


def element(parent,tag,attrs=None):
    e=OxmlElement(tag)
    for k,v in (attrs or {}).items():e.set(qn(k),str(v))
    parent.append(e);return e


def direction(p,rtl=True):
    pr=p._p.get_or_add_pPr()
    for c in list(pr.findall(qn('w:bidi'))):pr.remove(c)
    element(pr,'w:bidi',{'w:val':'1' if rtl else '0'})


def font(r,size=14,bold=False,latin=False):
    r.font.name=EN if latin else FA;r.font.size=Pt(size);r.font.bold=bold;r.font.italic=False
    r.font.color.rgb=RGBColor(0,0,0)
    pr=r._r.get_or_add_rPr();f=pr.rFonts
    for key in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+key),EN if latin else FA)
    element(pr,'w:szCs',{'w:val':str(round(size*2))})
    element(pr,'w:rtl',{'w:val':'0' if latin else '1'})
    element(pr,'w:lang',{'w:val':'en-US' if latin else 'fa-IR','w:bidi':'en-US' if latin else 'fa-IR'})
    element(pr,'w:bCs',{'w:val':'1' if bold else '0'})


def plain_runs(p,text,size=14,bold=False):
    # The supplied Nazanin lacks decimal/percent glyphs. Use the single Latin
    # face for complete numeric expressions instead of a fallback third font.
    digits=str.maketrans('۰۱۲۳۴۵۶۷۸۹٫٪','0123456789.%')
    text=re.sub(r'[۰-۹]+(?:٫[۰-۹]+)+(?:٪)?|[۰-۹]+٪',lambda m:m[0].translate(digits),text)
    for piece in re.split(r'([\u0600-\u06ff\u200c]+)',text):
        if piece:font(p.add_run(piece),size if re.search('[\u0600-\u06ff]',piece) else size*.86,bold,not re.search('[\u0600-\u06ff]',piece))


def bookmark(p,name):
    global BOOKMARK
    BOOKMARK+=1
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(BOOKMARK));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(BOOKMARK))
    p._p.insert(0,a);p._p.append(b)


def link(p,text,anchor=None,url=None,size=14):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),anchor)
    if url:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    from docx.text.paragraph import Paragraph
    if url:
        # Display-only break opportunities keep long repository URLs in the
        # margins; the actual hyperlink target remains the exact original URL.
        text=re.sub(r'([/_.-])',lambda m:m[1]+'\u200b',text)
    holder=OxmlElement('w:p');tmp=Paragraph(holder,p._parent);plain_runs(tmp,text,size)
    for run in list(holder):h.append(run)
    p._p.append(h)


def runs(p,text,size=14,bold=False):
    # Keep each citation group LTR, with a link to its numbered entry.
    pos=0
    for m in re.finditer(r'\[\d+\](?:,\s*\[\d+\])*',text):
        plain_runs(p,text[pos:m.start()],size,bold)
        control=p.add_run('\u202a');font(control,size*.86,latin=True)
        for j,n in enumerate(re.findall(r'\d+',m[0])):
            if j:plain_runs(p,', ',size,bold)
            link(p,'['+n+']',anchor='ref_'+n,size=size)
        control=p.add_run('\u202c');font(control,size*.86,latin=True)
        pos=m.end()
    plain_runs(p,text[pos:],size,bold)


def para(doc,text='',size=14,bold=False,after=6,style=None,justify=True):
    p=doc.add_paragraph(style=style);direction(p)
    p.alignment=ALIGN.JUSTIFY if justify else ALIGN.RIGHT
    if not justify:p._p.get_or_add_pPr().find(qn('w:jc')).set(qn('w:val'),'start')
    p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.18
    p.paragraph_format.widow_control=True
    runs(p,text,size,bold);return p


def heading(doc,text,level=1,page=False,anchor=None):
    p=para(doc,text,18 if level==1 else 15.5,True,8,'Heading '+str(level),False)
    p.paragraph_format.space_before=Pt(10);p.paragraph_format.keep_with_next=True
    p.paragraph_format.page_break_before=page
    if anchor:bookmark(p,anchor)
    return p


def table(doc,block):
    converted=[REFS.convert(s) for s in block]
    first=converted[0].strip('|').split('|')[0].strip()
    before=len(doc.paragraphs)
    old.table(doc,converted)
    p=doc.paragraphs[before]
    bookmark(p,'table_'+str(old.TABLE_COUNTER))
    title=p.text;TABLES.append((title,'table_'+str(old.TABLE_COUNTER)))
    # Captions and table cells retain their appropriate right/center alignment.
    p._p.get_or_add_pPr().find(qn('w:jc')).set(qn('w:val'),'start')


def figure(doc,index):
    name,title,note=FIGURES[index]
    p=para(doc,after=3);p.alignment=ALIGN.CENTER;p.paragraph_format.keep_with_next=True
    r=p.add_run();r.add_picture(str(ROOT/'figures/thesis_revision'/(name+'.png')),width=Inches(6.15))
    for e in r._r.xpath('.//wp:docPr'):e.set('descr',title)
    p=para(doc,'شکل '+str(index+1).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))+' '+title,12,True,5,justify=False)
    bookmark(p,'figure_'+str(index+1));p.paragraph_format.keep_with_next=True
    para(doc,REFS.convert(note),12,after=10)


def mr(text):
    r=OxmlElement('m:r');pr=element(r,'m:rPr');element(pr,'m:sty',{'m:val':'p'})
    wp=element(r,'w:rPr');element(wp,'w:rFonts',{'w:ascii':EN,'w:hAnsi':EN,'w:cs':EN})
    element(wp,'w:sz',{'w:val':'24'});element(wp,'w:szCs',{'w:val':'24'})
    element(r,'m:t').text=text;return r


def sub(a,b):
    e=OxmlElement('m:sSub');element(e,'m:e').append(mr(a));element(e,'m:sub').append(mr(b));return e


def nary(char,lower,upper,body):
    e=OxmlElement('m:nary');pr=element(e,'m:naryPr')
    element(pr,'m:chr',{'m:val':char});element(pr,'m:limLoc',{'m:val':'undOvr'})
    element(pr,'m:grow',{'m:val':'0'})
    if not upper:element(pr,'m:supHide',{'m:val':'1'})
    element(e,'m:sub').append(mr(lower));element(e,'m:sup').append(mr(upper))
    x=element(e,'m:e')
    for item in body:x.append(mr(item) if isinstance(item,str) else item)
    return e


def equation(doc,index):
    p=para(doc,after=5);direction(p,False);p.alignment=ALIGN.CENTER
    mp=OxmlElement('m:oMathPara');pr=element(mp,'m:oMathParaPr');element(pr,'m:jc',{'m:val':'center'})
    om=element(mp,'m:oMath')
    if index==1:
        pieces=[sub('F','q'),'(B) = 1[',sub('S','G'),' ⊆ ',sub('S','A'),'] ',
                nary('∏','i = 1','k',['1[',sub('w','i'),'(B) ≥ ',sub('r','i'),']'])]
    else:
        pieces=[sub('B','min'),' = ',sub('max','i'),' [',
                nary('∑','j < i','',['min(',sub('c','j'),', ',sub('r','i'),' − 1)']),
                ' + ',nary('∑','j ≤ i','',['1(',sub('c','j'),' ≥ ',sub('r','i'),')']),']']
    for item in pieces:om.append(mr(item) if isinstance(item,str) else item)
    p._p.append(mp)
    p=para(doc,'رابطه‌ی '+str(index).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')),11.5,after=8,justify=False)
    bookmark(p,'equation_'+str(index))


def markdown(doc,path,anchor=None,english=False):
    lines=path.read_text().splitlines();i=0;code=False
    while i<len(lines):
        raw=lines[i].strip();i+=1
        if not raw:continue
        if raw.startswith('```'):code=not code;continue
        # Insert figures at the point where the relevant result has been explained.
        before={'## ۳ ۷ سنجه‌ها و تجمیع':[0],'## ۴ ۷ سهم چارچوب و پیوند با تشخیص':[1],
                '## ۵ ۴ مسیر نامزد تا متن خروجی':[2],'## ۶ ۵ سه گواه برای فهم محدودیت':[3],
                '## ۶ ۸ کنترل شدت تغییر و اصلاحیه‌ی اکتشافی':[4,5]}
        for f in before.get(raw,[]):figure(doc,f)
        line=REFS.convert(raw) if not code else raw
        if raw.startswith('# '):
            text=line[2:];a=anchor
            if raw.startswith('# پیوست الف'):a='appendixA'
            if raw.startswith('# پیوست ب'):a='appendixB'
            if raw.startswith('# پیوست پ'):a='appendixC'
            p=heading(doc,text,page=True,anchor=a)
            if english:direction(p,False);p.alignment=ALIGN.LEFT
        elif raw.startswith('## '):heading(doc,line[3:],2)
        elif raw.startswith('|'):
            block=[raw]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
            table(doc,block)
        elif raw.startswith('!['):equation(doc,1 if 'feasibility_equation' in raw else 2)
        elif code or english:
            p=doc.add_paragraph();direction(p,False);p.alignment=ALIGN.LEFT if code else ALIGN.JUSTIFY
            p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.15
            font(p.add_run(line),10.5 if code else 12,latin=True)
        else:para(doc,line)


def nav(doc,title,items,page_map):
    heading(doc,title,page=True)
    for text,anchor in items:
        p=para(doc,after=7,justify=False)
        link(p,text,anchor=anchor,size=13)
        number=str(page_map.get(text,0)).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5),__import__('docx.enum.text',fromlist=['WD_TAB_ALIGNMENT']).WD_TAB_ALIGNMENT.LEFT)
        plain_runs(p,'\t'+number,13)


def build(page_map):
    global REFS,BOOKMARK,TABLES
    REFS=References();BOOKMARK=1000;TABLES=[]
    old.para=para;old.runs=runs;old.TABLE_COUNTER=0
    doc=Document();s=doc.sections[0];s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.9);s.left_margin=s.right_margin=Cm(2.1);s.footer_distance=Cm(.7)
    s.different_first_page_header_footer=True;element(s._sectPr,'w:bidi')
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name=FA;st.font.size=Pt(14);st.font.italic=False;st.font.color.rgb=RGBColor(0,0,0)
        pr=st.element.get_or_add_rPr()
        element(pr,'w:rFonts',{'w:ascii':EN,'w:hAnsi':EN,'w:cs':FA,'w:eastAsia':EN})
        element(pr,'w:szCs',{'w:val':'28'});element(pr,'w:iCs',{'w:val':'0'})
    for child in list(doc.styles.element.iter(qn('w:pBdr'))):child.getparent().remove(child)
    for child in list(doc.settings.element.findall(qn('m:mathPr'))):doc.settings.element.remove(child)
    mathpr=element(doc.settings.element,'m:mathPr');element(mathpr,'m:mathFont',{'m:val':EN})
    identity=university_mark(doc,'configs/current_report_identity.json')
    for text,size,bold,after,style in [(identity['university_fa'],18,True,8,None),
       (identity['course_fa'],15,False,16,None),('ارزیابی کنترل‌شده‌ی قطعه‌بندی',24,True,6,'Title'),
       ('و تحلیل تشخیصی بازیابی شاهد',23,True,6,'Subtitle'),('در حافظه‌ی گفت‌وگویی',23,True,18,'Subtitle'),
       ('نسخه‌ی بازبینی‌شده‌ی پایان‌نامه',16,False,14,None),('پژوهشگر  '+identity['author_fa'],16,True,9,None),
       ('استاد راهنمای دانشگاهی  '+identity['university_supervisor_fa'],15,False,9,None),
       ('راهنمای آموزشی علوم اعصاب شناختی  دکتر آرش فکری',14,False,7,None),
       ('فردیت FARDIAT مجتمع فوق تخصصی مغز و شناخت',14,False,8,None)]:para(doc,text,size,bold,after,style,False)
    p=para(doc,after=8,justify=False);p.add_run().add_picture(str(ROOT/identity['educational_institute_logo']),width=Inches(.6))
    para(doc,'۲۸ اکتبر ۲۰۲۶',14,after=10,justify=False)
    para(doc,'github.com/yousofab/event-memory-retrieval',11,after=0,justify=False)
    doc.core_properties.title='ارزیابی کنترل‌شده‌ی قطعه‌بندی و تحلیل تشخیصی بازیابی شاهد در حافظه‌ی گفت‌وگویی'
    doc.core_properties.author=identity['author_fa']
    doc.core_properties.subject='Supervisor revision: IEEE references, justified paragraphs, figures, equations and tool disclosure'
    footer=s.footer.paragraphs[0];direction(footer);footer.alignment=ALIGN.RIGHT
    footer._p.get_or_add_pPr().find(qn('w:jc')).set(qn('w:val'),'start')
    plain_runs(footer,'پایان‌نامه‌ی کارشناسی  |  نسخه‌ی بازبینی‌شده  |  ',10)
    field=element(footer._p,'w:fldSimple',{'w:instr':'PAGE'})
    markdown(doc,FOLDER/'front_matter_fa.md',anchor='abstractFA')
    markdown(doc,FOLDER/'abstract_en.md',anchor='abstractEN',english=True)
    items=[((FOLDER/'chapters'/n).read_text().splitlines()[0][2:],'ch'+n[:2]) for n in CHAPTERS]
    items += [('منابع','references'),('پیوست الف نمادگذاری و شرط دقیق','appendixA'),
              ('پیوست ب نقشه‌ی بازتولید','appendixB'),('پیوست پ واژه‌نامه','appendixC')]
    items.insert(0,('فهرست شکل‌ها و جدول‌ها','visualLists'))
    nav(doc,'فهرست مطالب',items,page_map)
    # Reserve a page for lists; populate it after table captions have been collected.
    list_heading=heading(doc,'فهرست شکل‌ها و جدول‌ها',page=True,anchor='visualLists')
    list_placeholder=para(doc,'',after=0)
    for n in CHAPTERS:markdown(doc,FOLDER/'chapters'/n,anchor='ch'+n[:2])
    appendix_start=len(doc.paragraphs)
    markdown(doc,FOLDER/'appendices_fa.md')
    first_appendix=doc.paragraphs[appendix_start]._p
    # References are generated after all first-use citations have been registered.
    reference_start=len(doc.paragraphs)
    heading(doc,'منابع',page=True,anchor='references')
    rows=REFS.save()
    for r in rows:
        p=doc.add_paragraph();direction(p,False);p.alignment=ALIGN.JUSTIFY
        p.paragraph_format.left_indent=Cm(.7);p.paragraph_format.first_line_indent=Cm(-.7)
        p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=1.08
        bookmark(p,'ref_'+str(r['number']))
        font(p.add_run(f'[{r["number"]}] '+r['text'].split(' [Online]')[0]),11.5,latin=True)
        font(p.add_run(' [Online]. Available: '),11.5,latin=True);link(p,r['url'],url=r['url'],size=13.37)
    reference_nodes=[p._p for p in doc.paragraphs[reference_start:]]
    for node in reference_nodes:first_appendix.addprevious(node)
    list_items=[('شکل '+str(i+1).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))+' '+f[1],'figure_'+str(i+1)) for i,f in enumerate(FIGURES)]+TABLES
    after=list_placeholder._p
    for title,anchor in list_items:
        p=para(doc,'',size=12,after=5,justify=False);link(p,title,anchor=anchor,size=12)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5),__import__('docx.enum.text',fromlist=['WD_TAB_ALIGNMENT']).WD_TAB_ALIGNMENT.LEFT)
        plain_runs(p,'\t'+str(page_map.get(title,0)).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')),12)
        after.addnext(p._p);after=p._p
    out=ROOT/'deliverables/thesis_supervisor_revision_fa.docx';doc.save(out)
    manifest={'body_fonts':[FA,EN],'equation_font_size_pt':12,'figures':len(FIGURES),'tables':old.TABLE_COUNTER,
              'references':len(rows),'lists':list_items,'body_justified':True}
    (FOLDER/'revision_layout.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(out)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--page-map',type=Path);a=p.parse_args()
    build(json.loads(a.page_map.read_text()) if a.page_map else {})
