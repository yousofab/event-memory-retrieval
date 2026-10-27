"""Build selected Persian thesis chapters using the existing report style.

Report-only dependencies: python-docx and a privately installed B Nazanin font.
The primary academic identity and corrected supervisor name are read from the
current identity snapshot. Historical report builders retain their old snapshot.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from package_phase19 import para, runs
from package_phase17 import element, rtl
from report_identity import university_mark

ROOT = Path(__file__).resolve().parents[1]


def title(doc, text, level=1, page=False):
    p = para(doc, text, 18 if level == 1 else 15.5, True, 8,
             'Heading ' + str(level))
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page


def table(doc, block):
    rows = [[s.strip() for s in line.strip('|').split('|')]
            for line in block if not re.fullmatch(r'[|:\-\s]+', line)]
    n = len(rows[0])
    widths = {2: [4.75, 1.65], 3: [3.25, 1.55, 1.60],
              4: [0.60, 2.40, 2.25, 1.15]}.get(n, [6.4 / n] * n)
    if n == 2 and rows[0][0] == 'کد':
        widths = [0.65, 5.75]
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    t.autofit = False
    element(t._tbl.tblPr, 'w:bidiVisual')
    element(t._tbl.tblPr, 'w:tblLayout', {'w:type': 'fixed'})
    borders = element(t._tbl.tblPr, 'w:tblBorders')
    for side in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        element(borders, 'w:' + side, {'w:val': 'single', 'w:sz': '4', 'w:color': 'D9D9D9'})
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        tr = t.rows[i]._tr.get_or_add_trPr()
        element(tr, 'w:cantSplit')
        if i == 0:
            element(tr, 'w:tblHeader')
        for j, text in enumerate(row):
            c = cells[j]
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            element(pr, 'w:shd', {'w:fill': 'E9EEF2' if i == 0 else 'FFFFFF'})
            margins = element(pr, 'w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                element(margins, 'w:' + side, {'w:w': '95', 'w:type': 'dxa'})
            p = c.paragraphs[0]
            rtl(p)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            if (n == 3 and j > 0) or (n == 2 and j == 1 and rows[0][0] != 'کد'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, text, 12.5 if rows[0][0] != 'کد' else 10.5, i == 0)
    para(doc, '', 3, after=4)


def markdown(doc, path):
    lines = path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('# '):
            title(doc, line[2:], page=True)
        elif line.startswith('## '):
            title(doc, line[3:], level=2)
        elif line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            table(doc, block)
            continue
        else:
            para(doc, line, 14, after=7)
        i += 1


def build():
    doc = Document()
    s = doc.sections[0]
    # Preserve the A4 academic report direction established by this project.
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.9)
    s.left_margin = s.right_margin = Cm(2.1)
    s.footer_distance = Cm(0.7)
    element(s._sectPr, 'w:bidi')
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        st = doc.styles[name]
        st.font.name = 'B Nazanin'
        st.font.size = Pt(14)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.italic = False
        pr = st.element.get_or_add_rPr()
        for k in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            pr.rFonts.set(qn('w:' + k), 'B Nazanin')
        element(pr, 'w:szCs', {'w:val': '28'})
        element(pr, 'w:iCs', {'w:val': '0'})
        ppr = st.element.get_or_add_pPr()
        for tag in ['w:pBdr', 'w:shd']:
            for child in list(ppr.findall(qn(tag))):
                ppr.remove(child)
    identity = university_mark(doc, 'configs/current_report_identity.json')
    para(doc, identity['university_fa'], 18, True, 8)
    para(doc, identity['course_fa'], 15, after=17)
    para(doc, 'تحلیل تشخیصی بازیابی شاهد', 25, True, 7, 'Title')
    para(doc, 'در حافظه‌ی گفت‌وگویی', 23, True, 15, 'Subtitle')
    para(doc, 'فصل‌های منتخب پایان‌نامه برای بررسی استاد', 16, after=14)
    para(doc, 'بسته‌ی اول تدوین فصل‌های ۵ و ۶ و ۷', 15, after=18)
    para(doc, 'پژوهشگر  ' + identity['author_fa'], 16, True, 8)
    para(doc, 'استاد راهنمای دانشگاهی  ' + identity['university_supervisor_fa'], 15, after=8)
    para(doc, 'راهنمای آموزشی علوم اعصاب شناختی  دکتر آرش فکری', 14, after=7)
    para(doc, 'فردیت FARDIAT مجتمع فوق تخصصی مغز و شناخت', 14, after=7)
    p = para(doc, after=9)
    p.add_run().add_picture(str(ROOT / identity['educational_institute_logo']), width=Inches(.65))
    for drawing in p._p.xpath('.//wp:docPr'):
        drawing.set('descr', 'فردیت FARDIAT خاستگاه آموزش علوم اعصاب شناختی')
    para(doc, '۲۷ اکتبر ۲۰۲۶', 14, after=12)
    para(doc, 'github.com/yousofab/event-memory-retrieval', 10.5, after=0)
    doc.core_properties.author = identity['author_fa']
    doc.core_properties.title = 'تحلیل تشخیصی بازیابی شاهد در حافظه‌ی گفت‌وگویی'
    doc.core_properties.subject = 'Selected BSc thesis chapters 5 6 7 for supervisor review'
    footer = s.footer.paragraphs[0]
    rtl(footer)
    runs(footer, 'پایان‌نامه‌ی کارشناسی  |  بسته‌ی اول تدوین  |  ', 10)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    title(doc, 'خلاصه‌ی یافته‌های مرکزی', page=True)
    para(doc, 'این پژوهش بررسی می‌کند چه عواملی مانع تبدیل قطعه‌بندی حافظه به پوشش شاهد می‌شوند. یافته‌های موجود حول چارچوب ارزیابی کنترل‌شده، تجزیه‌ی دقیق شکست و تمایز دسترسی و مونتاژ شاهد سازمان یافته‌اند. اهمیت عملی آن، مشخص کردن محل و ظرفیت اصلاح در زنجیره‌ی بازیابی با بودجه‌ی محدود است.')
    para(doc, 'در فضای ثابت مرحله‌ی ۱۸، از ۱۵۶ شکست LoCoMo، تعداد ۱۰۶ مورد جلسه‌ی شاهد غایب دارند، ۲۵ مورد با هندسه‌ی موجود ناممکن‌اند و ۲۵ مورد فرصت موفقِ انتخاب‌نشده دارند. سقف مشروط ۴۶۸ از ۵۹۹ است. مرحله‌ی ۱۹ نشان می‌دهد افزایش دسترسی به جلسه می‌تواند با افت پوشش شاهد همراه شود. QMSum نیز هم‌راستایی تقریبی مرز موضوع را با بهبود بازیابی همراه نشان نداده است.')
    title(doc, 'مسیر مطالعه‌ی این بسته', level=2)
    para(doc, 'فصل ۵ شکاف هم‌راستایی مرز و بازیابی را تحلیل می‌کند. فصل ۶ تعریف ممیز، تجزیه‌ی شکست و مداخله‌ی دسترسی را توضیح می‌دهد. فصل ۷ اهمیت، کاربرد، محدودیت‌ها و مسیرهای آینده را جمع‌بندی می‌کند. راهنمای انتهایی، ارجاع کدهای منبع و پرسش‌های بازبینی را مشخص می‌کند.')
    para(doc, 'اعداد از خروجی‌های ثبت‌شده خوانده شده‌اند. شناسه‌های C به دفتر شواهد عددی و شناسه‌های S به منابع داخلی پژوهش ارجاع دارند. متن حاضر برای بررسی فصل‌های منتخب تهیه شده است؛ تدوین یکپارچه‌ی همه‌ی فصل‌ها پس از بازخورد انجام می‌شود.')
    for name in ['05_qmsum_gap_fa.md', '06_diagnostics_fa.md', '07_discussion_fa.md']:
        markdown(doc, ROOT / 'thesis/chapters' / name)
    markdown(doc, ROOT / 'thesis/review_guide_fa.md')
    out = ROOT / 'deliverables/thesis_batch1_fa.docx'
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build()
