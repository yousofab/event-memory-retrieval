"""Build editable RTL B Nazanin Phase 12 research report."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_phase9_report as prior


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/phase12.md"
DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۱۲-حافظه-و-اصلاح-اطلاعات.docx"


def build():
    prior.base.SOURCE = SOURCE
    prior.base.DESTINATION = DESTINATION
    prior.build()
    doc = Document(DESTINATION)
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 1.08
        if paragraph.text.startswith(("[۱]", "[۲]", "[۳]", "[۴]", "[۵]")):
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
    footer = section.footer.paragraphs[0]
    footer.clear()
    prior.base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    prior.base.add_run(footer, "گزارش پژوهشی مرحله‌ی دوازدهم  |  صفحه ", size=10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "گزارش پژوهشی مرحله‌ی دوازدهم حافظه و اصلاح اطلاعات"
    doc.core_properties.subject = "دسترسی به شواهد به‌روزشده در حافظه‌ی گفت‌وگو"
    doc.save(DESTINATION)
    return DESTINATION


if __name__ == "__main__":
    print(build())
