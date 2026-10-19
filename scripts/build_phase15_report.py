"""Build the Persian RTL Phase 15 Word report from its source Markdown."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import build_phase14_report as prior


ROOT = Path(__file__).resolve().parents[1]
prior.SOURCE = ROOT / "docs/phase15.md"
prior.DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۱۵-نقش-و-هزینه-Memora.docx"


def build():
    prior.build()
    doc = Document(prior.DESTINATION)
    doc.sections[0].top_margin = Cm(1.6)
    doc.sections[0].bottom_margin = Cm(1.6)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 1.02
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                if run.font.name == "B Nazanin" and run.font.size and run.font.size.pt == 12.5:
                    run.font.size = Pt(11.8)
        elif paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    prior.prior.base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    prior.prior.base.add_run(footer, "گزارش پژوهشی مرحله‌ی پانزدهم  |  صفحه ", size=10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "گزارش پژوهشی مرحله‌ی پانزدهم نقش پیام کاربر و هزینه‌ی متن"
    doc.core_properties.subject = "بازیابی شواهد در Memora با کنترل نقش و بودجه"
    doc.save(prior.DESTINATION)
    return prior.DESTINATION


if __name__ == "__main__":
    print(build())
