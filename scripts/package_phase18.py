"""Build the Persian academic report; requires python-docx and B Nazanin."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from package_phase17 import element, rtl, runfont

ROOT = Path(__file__).resolve().parents[1]
FONT = 'B Nazanin'


def runs(p, text, size=14, bold=False, color='000000'):
    text = text.replace('**', '').replace('`', '')
    for part in re.split(r'([\x21-\x7e]+(?:[ \t]+[\x21-\x7e]+)*)', text):
        if part:
            runfont(p.add_run(part), size, bold, bool(re.fullmatch(r'[\x20-\x7e]+', part)), color)


def para(doc, text='', size=14, bold=False, after=6, style=None):
    p = doc.add_paragraph(style=style)
    rtl(p)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    runs(p, text, size, bold)
    return p


def heading(doc, text, page=False):
    p = para(doc, text, 17, True, 8, 'Heading 1')
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.page_break_before = page
    return p


def table(doc, block, number):
    headings = {1:'داده و معنای برچسب', 2:'تفکیک نتیجه‌ی نویسنده در ۱۰۲۴ واژه',
                3:'شمار پوشش کامل در سه بودجه', 4:'تفاوت سطح معیار', 5:'آمار کمترین بودجه‌ی متناهی',
                6:'نمونه‌های توصیفی از گواه‌های عددی'}
    p = para(doc, f'جدول {number} '+headings[number], 11.5, True, 5)
    p.paragraph_format.keep_with_next = True
    rows = [[c.strip() for c in line.strip('|').split('|')] for line in block
            if not re.fullmatch(r'[|:\-\s]+', line)]
    n = len(rows[0])
    widths = [3.25, 1.64, 1.64] if n == 3 else [1.05, 1.35, 1.38, 1.35, 1.40]
    if number == 1:
        widths = [1.05, 2.25, 3.23]
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    t.autofit = False
    pr = t._tbl.tblPr
    element(pr, 'w:bidiVisual')
    element(pr, 'w:tblLayout', {'w:type':'fixed'})
    borders = element(pr, 'w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        element(borders, 'w:'+side, {'w:val':'single','w:sz':'4','w:color':'D9D9D9'})
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        trpr = t.rows[i]._tr.get_or_add_trPr()
        element(trpr, 'w:cantSplit')
        if i == 0:
            element(trpr, 'w:tblHeader')
        for j, value in enumerate(row):
            c = cells[j]
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            element(tcpr, 'w:shd', {'w:fill':'264B65' if i == 0 else ('F3F6F8' if i % 2 else 'FFFFFF')})
            margin = element(tcpr, 'w:tcMar')
            for side in ('top','bottom','left','right'):
                element(margin, 'w:'+side, {'w:w':'85','w:type':'dxa'})
            p = c.paragraphs[0]
            rtl(p)
            if (number == 1 and j < 2) or (number != 1 and j > 0):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.04
            p.paragraph_format.space_after = Pt(2)
            runs(p, value, 12, i == 0, 'FFFFFF' if i == 0 else '000000')
    para(doc, '', 3, after=3)


def build():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(1.9), Cm(1.85)
    s.left_margin = s.right_margin = Cm(2.2)
    s.footer_distance = Cm(.8)
    element(s._sectPr, 'w:bidi')
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(14)
        st.font.color.rgb = RGBColor(0,0,0)
        st.font.italic = False
        pr = st.element.get_or_add_rPr()
        for key in ('ascii','hAnsi','cs','eastAsia'):
            pr.rFonts.set(qn('w:'+key), FONT)
        element(pr, 'w:szCs', {'w:val':'28'})
        element(pr, 'w:iCs', {'w:val':'0'})
        ppr = st.element.get_or_add_pPr()
        for tag in ('w:pBdr','w:shd'):
            for child in list(ppr.findall(qn(tag))):
                ppr.remove(child)
    doc.core_properties.author = 'Yousef Abolfathi Roozbahani'
    doc.core_properties.title = 'امکان‌پذیری بازیابی شاهد در پنجره‌های ثابت'
    doc.core_properties.subject = 'Phase 18 exact feasibility and minimum-budget audit'
    footer = s.footer.paragraphs[0]
    rtl(footer)
    runs(footer, 'پژوهش حافظه‌ی رویدادی  |  مرحله‌ی ۱۸  |  ', 10)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)

    para(doc, 'دانشگاه آزاد اسلامی واحد کرج', 18, True, 14)
    para(doc, 'کارشناسی مهندسی کامپیوتر  |  پروژه‌ی نرم‌افزار', 16, after=38)
    para(doc, 'گزارش پژوهشی مرحله ۱۸', 26, True, 16, 'Title')
    para(doc, 'امکان‌پذیری بازیابی شاهد در پنجره‌های ثابت', 23, True, 18, 'Subtitle')
    para(doc, 'تحلیل شکست و محاسبه‌ی کمترین بودجه با گواه قابل بررسی', 16, after=34)
    para(doc, 'پژوهشگر  یوسف ابوالفتحی روزبهانی', 16, True, 14)
    para(doc, 'استاد راهنمای دانشگاهی  استاد امیر نوری ولنده', 15, after=10)
    para(doc, 'راهنمای آموزشی علوم اعصاب شناختی  دکتر آرش فکری', 15, after=10)
    para(doc, 'فردیت  FARDIAT  مجتمع فوق تخصصی مغز و شناخت', 15, after=18)
    logo = ROOT / 'assets/fardiat-logo.jpg'
    if logo.exists():
        p = para(doc, after=12)
        p.add_run().add_picture(str(logo), width=Inches(1.05))
    para(doc, '۲۴ اکتبر ۲۰۲۶', 14, after=12)
    para(doc, 'گزارش پیشرفت پژوهش', 13, after=8)
    para(doc, 'github.com/yousofab/event-memory-retrieval', 11, after=0)

    parts = re.split(r'^## (.+)$', (ROOT / 'docs/phase18.md').read_text(), flags=re.M)
    table_no = 0
    for title, body in zip(parts[1::2], parts[2::2]):
        page = title == 'چکیده' or title.startswith(('۲ ','۳ ','۵ ','۶ ','۷ ','۸ ','۱۰ '))
        heading(doc, title, page)
        lines = body.strip().splitlines()
        i, code = 0, False
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            if line.startswith('```'):
                code = not code
                i += 1
                continue
            if code:
                p = doc.add_paragraph()
                rtl(p, False)
                runfont(p.add_run(line), 9.0, latin=True, color='000000')
                p.paragraph_format.space_after = Pt(4)
                i += 1
                continue
            if line.startswith('|'):
                block = []
                while i < len(lines) and lines[i].startswith('|'):
                    block.append(lines[i]); i += 1
                table_no += 1
                table(doc, block, table_no)
                continue
            if line.startswith('!['):
                m = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
                p = para(doc, after=0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str((ROOT / 'docs' / m[2]).resolve()), width=Inches(6.1))
                para(doc, m[1], 11.5, after=6)
                i += 1
                continue
            para(doc, line)
            i += 1

    refs = json.loads((ROOT / 'docs/phase18_literature_atlas.json').read_text())
    for start, stop in ((0,9),(9,17)):
        heading(doc, 'منابع' if not start else 'ادامه‌ی منابع', True)
        if not start:
            para(doc, 'شناسه‌های S01 تا S17 با اطلس همراه یکسان‌اند. شناسنامه‌ی کامل، دامنه‌ی مطالعه، نسخه و هش منابع در اطلس ثبت شده است. دسترسی ۲۴ اکتبر ۲۰۲۶.', 11.5, after=8)
        for r in refs[start:stop]:
            names = []
            for a in r['authors']:
                a = re.sub(r'\s*\([^)]*[\u4e00-\u9fff][^)]*\)', '', a).strip()
                if a.count(',') == 1:
                    family, given = a.split(','); a = given.strip()+' '+family.strip()
                names.append(a)
            if len(names) > 6:
                names = names[:3] + ['et al.']
            p = doc.add_paragraph()
            rtl(p, False)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.02
            runfont(p.add_run(f"[{r['id']}] {', '.join(names)}. {r['title']}."), 10.5, latin=True, color='000000')
            p = para(doc, r['venue_and_status']+'؛ '+str(r['date'])+(('؛ '+r['arxiv_version']) if r['arxiv_version'] else ''), 11.5, after=2)
            p.paragraph_format.keep_with_next = True
            p = doc.add_paragraph()
            rtl(p, False)
            p.paragraph_format.space_after = Pt(8)
            runfont(p.add_run(r['reference_url']), 9, latin=True, color='000000')
    out = ROOT / 'deliverables/phase18_research_report_fa.docx'
    out.parent.mkdir(exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build()
