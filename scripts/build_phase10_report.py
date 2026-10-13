"""Build the Phase 10 Persian RTL B Nazanin report with the supplied mark."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
import build_phase9_report as prior


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/phase10.md"
DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۱۰-مرزهای-موضوعی.docx"


def build():
    prior.base.SOURCE = SOURCE
    prior.base.DESTINATION = DESTINATION
    prior.build()
    document = Document(DESTINATION)
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 1.10
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    prior.base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    prior.base.add_run(footer, "گزارش پژوهشی مرحله‌ی دهم  |  صفحه ", size=10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "گزارش پژوهشی مرحله‌ی دهم سنجش مرزهای موضوعی"
    document.core_properties.subject = "انتقال مرزبند علّی به برچسب‌های موضوعی QMSum"
    document.save(DESTINATION)
    return DESTINATION


if __name__ == "__main__":
    print(build())
