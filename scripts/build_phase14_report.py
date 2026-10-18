"""Generate the editable Phase 14 Persian RTL report from docs/phase14.md."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_phase9_report as prior


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/phase14.md"
DESTINATION = ROOT / "deliverables/گزارش-پژوهشی-مرحله-۱۴-ممیزی-ورودی-Memora.docx"


def build():
    prior.base.SOURCE = SOURCE
    prior.base.DESTINATION = DESTINATION
    prior.build()
    doc = Document(DESTINATION)
    section = doc.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 1.09
        if paragraph.text == "سپاسگزاری":
            logo = ROOT / "assets/fardiat-logo.jpg"
            if logo.exists():
                figure = doc.add_paragraph()
                prior.base.persian_paragraph(figure, WD_ALIGN_PARAGRAPH.CENTER)
                figure.paragraph_format.space_after = Pt(4)
                figure.add_run().add_picture(str(logo), width=Cm(2.6))
                paragraph._p.addnext(figure._p)
    footer = section.footer.paragraphs[0]
    footer.clear()
    prior.base.persian_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    prior.base.add_run(footer, "گزارش پژوهشی مرحله‌ی چهاردهم  |  صفحه ", size=10)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    doc.core_properties.title = "گزارش پژوهشی مرحله‌ی چهاردهم ممیزی وفاداری ورودی Memora"
    doc.core_properties.subject = "ممیزی نقش گوینده در مسیر ارزیابی مستقیم مدل"
    # Replace raw markdown-style references with compact URLs that are live in Word.
    in_refs = False
    for paragraph in list(doc.paragraphs):
        if paragraph.text == "منابع":
            in_refs = True
            continue
        if in_refs and paragraph.text.startswith("["):
            line = paragraph.text
            if " https://" in line:
                citation, url = line.rsplit(" https://", 1)
                paragraph.clear()
                prior.base.persian_paragraph(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
                paragraph.paragraph_format.space_after = Pt(5)
                run = paragraph.add_run(citation + "  ")
                run.font.name = "DejaVu Sans"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0, 0, 0)
                prior.add_hyperlink(paragraph, "https://" + url)
    doc.save(DESTINATION)
    return DESTINATION


if __name__ == "__main__":
    print(build())
