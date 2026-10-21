"""Build the Persian Phase17 report from audited project text and results.

Requires python-docx; B Nazanin must be installed for faithful rendering.
The experiment itself has no document or plotting dependency.
"""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
FONT='B Nazanin'; LATIN='Calibri'; INK='153C4B'; TEAL='096B73'

def element(parent,tag,attrs=None):
    e=OxmlElement(tag)
    for k,v in (attrs or {}).items():e.set(qn(k),str(v))
    parent.append(e);return e

def rtl(p,value=True):
    pr=p._p.get_or_add_pPr();element(pr,'w:bidi',{'w:val':'1' if value else '0'})
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    # Logical start resolves to the physical right in RTL in Word/LibreOffice.
    # A literal 'right' is mirrored by this renderer for bidi paragraphs.
    if value:pr.find(qn('w:jc')).set(qn('w:val'),'start')

def runfont(run,size=14,bold=False,latin=False,color=None):
    run.font.name=LATIN if latin else FONT;run.font.size=Pt(size);run.font.bold=bold
    pr=run._r.get_or_add_rPr();fonts=pr.rFonts
    for name in ['ascii','hAnsi','cs','eastAsia']:fonts.set(qn('w:'+name),LATIN if latin else FONT)
    element(pr,'w:szCs',{'w:val':int(size*2)})
    element(pr,'w:rtl',{'w:val':'0' if latin else '1'})
    element(pr,'w:lang',{'w:val':'en-US' if latin else 'fa-IR','w:bidi':'fa-IR'})
    if bold:element(pr,'w:bCs')
    if color:run.font.color.rgb=RGBColor.from_string(color)

def clean(text):
    return text.replace('**','').replace('`','')

def text_runs(p,text,size=14,bold=False,color=None):
    # English runs remain LTR inside a genuinely RTL paragraph.
    for part in re.split(r'([A-Za-z][A-Za-z0-9_/.@:+=;,%()− -]*[A-Za-z0-9_/.]|[A-Za-z])',clean(text)):
        if part:runfont(p.add_run(part),size,bold,bool(re.match('[A-Za-z]',part)),color)

def para(doc,text='',size=14,bold=False,after=5,color=None):
    p=doc.add_paragraph();rtl(p);p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.12
    text_runs(p,text,size,bold,color);return p

def head(doc,text,level=1):
    p=doc.add_paragraph(style=f'Heading {level}');rtl(p)
    p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(7)
    text_runs(p,text,17 if level==1 else 15,True,INK);return p

def table(doc,lines,n):
    para(doc,f'جدول {n}. '+{1:'داده و سنجه‌ی ارزیابی',2:'نتیجه‌ی اصلی در ۱۰۲۴ واژه',3:'اختلاف‌های جفت‌شده و عدم قطعیت',4:'حساسیت به بودجه'}.get(n,'مقایسه‌ی روش‌ها'),12,True,4,TEAL)
    rows=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines if not re.match(r'^\|[ :\-]+\|',l)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.RIGHT;t.autofit=False
    element(t._tbl.tblPr,'w:bidiVisual')
    element(t._tbl.tblPr,'w:tblLayout',{'w:type':'fixed'})
    total=6.53
    widths=([2.25,1.42,1.42,1.44] if len(rows[0])==4 else [2.5,1.8,2.23])
    for j,c in enumerate(t.columns):c.width=Inches(widths[j])
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=t.rows[i]._tr.get_or_add_trPr();element(trpr,'w:cantSplit')
        if i==0:element(trpr,'w:tblHeader')
        for j,value in enumerate(row):
            c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();element(tcpr,'w:shd',{'w:fill':'E8F1F2' if i==0 else ('F7F9FA' if i%2 else 'FFFFFF')})
            mar=element(tcpr,'w:tcMar')
            for side in ['top','bottom','left','right']:element(mar,'w:'+side,{'w:w':'75','w:type':'dxa'})
            p=c.paragraphs[0];rtl(p);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
            text_runs(p,value,12.5,i==0,INK if i==0 else None)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def build():
    doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.85);s.bottom_margin=Cm(1.85);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2);s.footer_distance=Cm(.9)
    element(s._sectPr,'w:bidi')
    for name in ['Normal','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name=FONT;st.font.size=Pt(14)
        pr=st.element.get_or_add_rPr();fonts=pr.rFonts
        for x in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+x),FONT)
        element(pr,'w:szCs',{'w:val':'28'})
    footer=s.footer.paragraphs[0];rtl(footer);text_runs(footer,'پژوهش حافظه‌ی رویدادی  |  مرحله‌ی ۱۷  |  ',10,color='667784')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.core_properties.author='Yousef Abolfathi Roozbahani';doc.core_properties.title='مرحله‌ی ۱۷: ارزیابی اثر مرز با فرصت ثابت';doc.core_properties.subject='Event memory evidence retrieval; controlled exploratory study'
    para(doc,'دانشگاه آزاد اسلامی واحد کرج',17,True,14,INK)
    para(doc,'کارشناسی مهندسی کامپیوتر | پروژه‌ی نرم‌افزار',15,False,45)
    para(doc,'گزارش پژوهشی مرحله‌ی ۱۷',24,True,18,INK)
    para(doc,'مرز رویداد با فرصت و هزینه‌ی بازیابی ثابت',23,True,18,TEAL)
    para(doc,'ممیزی علمی، مرور مرز دانش و اجرای آزمایش کنترل‌شده',16,False,40)
    para(doc,'پژوهشگر: یوسف ابوالفتحی روزبهانی',16,True,14)
    para(doc,'استاد راهنمای دانشگاهی: استاد امیر نوری ولنده',15,False,10)
    para(doc,'راهنمای آموزشی علوم اعصاب شناختی: دکتر آرش فکری',15,False,10)
    para(doc,'FARDIAT | فردیت، مجتمع فوق تخصصی مغز و شناخت',15,False,24)
    logo=ROOT/'assets/fardiat-logo.jpg'
    if logo.exists():
        p=doc.add_paragraph();rtl(p);p.add_run().add_picture(str(logo),width=Inches(1.05))
    para(doc,'۲۱ اکتبر ۲۰۲۶',14,False,12)
    para(doc,'گزارش پیشرفت برای بازبینی؛ انتشار عمومی پس از تایید پژوهشگر',12,False,10,'667784')
    para(doc,'github.com/yousofab/event-memory-retrieval',12,False,2,TEAL)
    raw=(ROOT/'docs/phase17.md').read_text();parts=re.split(r'^## ',raw,flags=re.M)[1:]
    sections={p.split('\n',1)[0]:p.split('\n',1)[1] for p in parts}
    table_no=0
    for idx,(title,body) in enumerate(sections.items()):
        if title=='پیوند اجزای بسته':continue
        # Combine abstract/necessity, historical/audit, and limitations/thanks.
        h=head(doc,title)
        if not(title.startswith(('۱.','۸.','۱۰.'))):h.paragraph_format.page_break_before=True
        lines=body.strip().splitlines();i=0
        while i<len(lines):
            line=lines[i].strip()
            if not line:i+=1;continue
            if line.startswith('|'):
                block=[]
                while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
                table_no+=1;table(doc,block,table_no);continue
            if line.startswith('!['):
                m=re.match(r'!\[(.*?)\]\((.*?)\)',line);path=(ROOT/'docs'/m[2]).resolve()
                p=doc.add_paragraph();rtl(p)
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(path),width=Inches(5.85))
                para(doc,m[1]+'.',11,False,5,'52646B')
                i+=1;continue
            para(doc,line,14,False,6);i+=1
    head(doc,'راهنمای بسته و بازتولید').paragraph_format.page_break_before=True
    para(doc,'تمام نتیجه‌ها در فایل‌های عددی قابل بررسی‌اند. کد اصلی فقط Python استاندارد و CPU می‌خواهد. متن خام داده، قلم و مقاله‌های دانلودشده در بسته نیستند. برای حفظ سابقه، خروجی بازاجرا در مسیر تازه نوشته می‌شود.')
    for label,path in [('اجرای اصلی','scripts/run_phase17_fixed_anchors.py'),('کنترل برابری و اعداد','scripts/validate_phase17.py'),('تحلیل مکمل و شکل‌ها','scripts/analyze_phase17.py'),('خروجی تفصیلی','results/phase17_rows.json.gz'),('خلاصه و عدم قطعیت','results/phase17_summary.json'),('گواه بازاجرا و آزمون‌ها','results/phase17_validation.json'),('اطلس ۳۷ منبع و شناسنامه','docs/phase17_literature_atlas.md'),('ممیزی مراحل ۱ تا ۱۶','docs/phase17_state_ledger.md'),('ده ایده و نقد سه گزینه','docs/phase17_candidates.md'),('تحلیل شکاف و پیشینه','docs/phase17_frontier_synthesis.md'),('راهنمای کامل اجرا','docs/phase17_reproduce.md')]:
        para(doc,label+' — '+path,12.5,False,5)
    head(doc,'فرمان‌های اجرای اصلی',2)
    for line in ['python scripts/fetch_data.py','python scripts/fetch_phase10_data.py --download','python -m unittest discover -s tests -v','python scripts/run_phase17_fixed_anchors.py --output-dir /tmp/phase17-review-run','python scripts/validate_phase17.py --directory /tmp/phase17-review-run']:
        p=doc.add_paragraph();rtl(p,False);runfont(p.add_run(line),9,latin=True);p.paragraph_format.space_after=Pt(5)
    para(doc,'شاخه‌ی آماده: research/phase17-fixed-anchor-evidence. متن PR و بسته‌ی تغییرات آماده‌اند. هیچ push، PR عمومی یا release در این چرخه انجام نشده است.',12.5,False,6)
    refs=json.loads((ROOT/'docs/phase17_literature_atlas.json').read_text())
    for start in range(0,len(refs),8):
        head(doc,'منابع و شناسنامه‌ی انتشار'+('؛ ادامه' if start else '')).paragraph_format.page_break_before=True
        if start==0:para(doc,'شناسه‌های R01 تا R37 با اطلس همراه یکسان‌اند. ارجاع‌ها به سبک عددی تنظیم شده‌اند؛ اطلاعات کامل نویسندگان، دامنه‌ی مطالعه، کد، رتبه‌ی مجله و هش منبع در اطلس ماشینی محفوظ است. دسترسی: ۲۱ اکتبر ۲۰۲۶.',11.5,False,7)
        for r in refs[start:start+8]:
            authors=[]
            for author in r['authors']:
                a=re.sub(r'\s*\([^)]*[\u4e00-\u9fff][^)]*\)','',author).strip()
                if a.count(',')==1:
                    family,given=a.split(',');a=given.strip()+' '+family.strip()
                authors.append(a)
            by=', '.join(authors if len(authors)<=6 else authors[:3]+['et al'])
            p=doc.add_paragraph();rtl(p,False);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.03
            text=f"[{r['id']}] {by}. {r['title']}."
            runfont(p.add_run(text),10.5,latin=True)
            para(doc,r['venue_and_status']+'؛ '+r['reading_scope']+'.',11.5,False,2)
            url=('https://doi.org/'+r['doi']) if r['doi'] else r['publication_url']
            p=doc.add_paragraph();rtl(p,False);p.paragraph_format.space_after=Pt(9)
            runfont(p.add_run(url),9,latin=True,color=TEAL)
    out=ROOT/'deliverables/phase17_research_report_fa.docx';doc.save(out);print(out)

if __name__=='__main__':build()
