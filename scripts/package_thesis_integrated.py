"""Build the integrated Persian review manuscript from its eight source chapters.

Uses the project's A4 RTL B Nazanin academic identity. A rendered-PDF outline
provides the static, clickable contents page; pass its JSON with --page-map.
"""
import argparse
import json
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from package_phase17 import element, rtl, runfont
from package_phase19 import para, runs
from report_identity import university_mark
from build_thesis_synthesis import CHAPTERS

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT/'thesis'
TABLE_COUNTER = 0
EQUATION_COUNTER = 0


def bookmark(p, name, number):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(number));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(number))
    p._p.insert(0,a);p._p.append(b)


def link(p, text, anchor=None, url=None, size=14):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),anchor)
    if url:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    # Apply the same carefully split RTL/LTR runs used in ordinary paragraphs.
    temporary=p.add_run();p._p.remove(temporary._r)
    from docx.text.paragraph import Paragraph
    holder=OxmlElement('w:p');tmp=Paragraph(holder,p._parent)
    runs(tmp,text,size)
    for r in list(holder):h.append(r)
    p._p.append(h)


def heading(doc, text, level=1, page=False, anchor=None):
    p=para(doc,text,18 if level==1 else 15.5,True,8,'Heading '+str(level))
    p.paragraph_format.space_before=Pt(9)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.page_break_before=page
    if anchor:bookmark(p,anchor,100+len(doc.paragraphs))
    return p


def table(doc, block):
    global TABLE_COUNTER
    rows=[[s.strip() for s in line.strip('|').split('|')]
          for line in block if not re.fullmatch(r'[|:\-\s]+',line)]
    n=len(rows[0]);first=rows[0][0]
    widths=([1.5,1.5,1.5,1.9] if n==4 else [1.35,2.5,2.55] if n==3 else [3.9,2.5])
    if first=='کد':widths=[.65,5.75]
    if first=='اصطلاح':widths=[1.45,4.95]
    if first=='سنجه':widths=[3.25,1.55,1.60]
    if first=='وضعیت در ۱۰۲۴ واژه':widths=[4.75,1.65]
    if first.startswith('شاخص LoCoMo'):widths=[3.3,1.55,1.55]
    captions={'مطالعه':'مقایسه‌ی طراحی‌های نزدیک','منبع':'داده‌ها و معنای مرجع',
              'جزء':'اجزای مشترک و متغیر طراحی','سنجه':'مرز موضوعی و پوشش بازیابی',
              'وضعیت در ۱۰۲۴ واژه':'تجزیه‌ی پرسش‌های LoCoMo',
              'کد':'منابع داخلی پژوهش','شناسه':'ادعاهای تکمیلی','اصطلاح':'واژه‌نامه'}
    TABLE_COUNTER += 1
    caption='دسترسی و امکان و پوشش' if first.startswith('شاخص LoCoMo') else captions.get(first,first)
    number=str(TABLE_COUNTER).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
    p=para(doc,'جدول '+number+' '+caption,12,True,5)
    p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.RIGHT
    element(t._tbl.tblPr,'w:bidiVisual');element(t._tbl.tblPr,'w:tblLayout',{'w:type':'fixed'})
    borders=element(t._tbl.tblPr,'w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        element(borders,'w:'+side,{'w:val':'single','w:sz':'4','w:color':'D9D9D9'})
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        pr=t.rows[i]._tr.get_or_add_trPr();element(pr,'w:cantSplit')
        if i==0:element(pr,'w:tblHeader')
        for j,text in enumerate(row):
            c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();element(pr,'w:shd',{'w:fill':'E9EEF2' if i==0 else 'FFFFFF'})
            margins=element(pr,'w:tcMar')
            for side in ['top','bottom','left','right']:element(margins,'w:'+side,{'w:w':'95','w:type':'dxa'})
            p=c.paragraphs[0];rtl(p);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=i==0
            if len(text)<22 and j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            runs(p,text,10.5 if first=='کد' else 12.5,i==0)
    para(doc,'',3,after=4)


def markdown(doc, path, anchor=None, english=False):
    global EQUATION_COUNTER
    lines=path.read_text().splitlines();i=0;code=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('```'):code=not code;i+=1;continue
        if line.startswith('# '):
            text=line[2:];a=anchor
            if text.startswith('پیوست الف'):a='appendixA'
            elif text.startswith('پیوست ب'):a='appendixB'
            elif text.startswith('پیوست پ'):a='appendixC'
            p=heading(doc,text,page=True,anchor=a)
            if english:rtl(p,False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('## '):
            p=heading(doc,line[3:],level=2)
            if re.fullmatch(r'P\d{2}',line[3:]):bookmark(p,line[3:],100+len(doc.paragraphs))
        elif line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip());i+=1
            table(doc,block);continue
        elif re.fullmatch(r'!\[.*\]\(.*\)',line):
            match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
            p=para(doc,after=3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            r=p.add_run();r.add_picture(str((path.parent/match[2]).resolve()),width=Inches(6.2))
            for d in r._r.xpath('.//wp:docPr'):d.set('descr',match[1])
            EQUATION_COUNTER += 1
            number=str(EQUATION_COUNTER).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
            para(doc,'رابطه‌ی '+number+' '+match[1],11.5,after=8)
        elif line.startswith('https://'):
            p=doc.add_paragraph();rtl(p,False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after=Pt(8);link(p,line,url=line,size=9)
        elif code or english or (path.name=='references_fa.md' and re.match(r'[A-Z]',line)):
            p=doc.add_paragraph();rtl(p,False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.12
            runfont(p.add_run(line),10 if code or path.name=='references_fa.md' else 11,latin=True,color='000000')
        else:para(doc,line,14,after=6)
        i+=1


def contents(doc, page_map):
    heading(doc,'فهرست مطالب',page=True)
    items=[((FOLDER/'chapters'/n).read_text().splitlines()[0][2:],'ch'+n[:2]) for n in CHAPTERS]
    items += [('منابع بیرونی','externalSources'),('منابع داخلی و راهنمای شواهد','internalSources'),
              ('پیوست الف نمادگذاری و شرط دقیق','appendixA'),
              ('پیوست ب نقشه‌ی بازتولید','appendixB'),('پیوست پ واژه‌نامه','appendixC')]
    t=doc.add_table(rows=0,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.RIGHT
    element(t._tbl.tblPr,'w:bidiVisual');element(t._tbl.tblPr,'w:tblLayout',{'w:type':'fixed'})
    t.columns[0].width=Inches(5.7);t.columns[1].width=Inches(.7)
    for title,anchor in items:
        cells=t.add_row().cells;cells[0].width=Inches(5.7);cells[1].width=Inches(.7)
        p=cells[0].paragraphs[0];rtl(p);p.paragraph_format.space_after=Pt(9)
        link(p,title,anchor=anchor)
        p=cells[1].paragraphs[0];rtl(p);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        number=str(page_map.get(title,0)).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
        runs(p,number,14)
    return items


def build(page_map):
    global TABLE_COUNTER, EQUATION_COUNTER
    TABLE_COUNTER=EQUATION_COUNTER=0
    doc=Document();s=doc.sections[0]
    s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(1.9);s.left_margin=s.right_margin=Cm(2.1);s.footer_distance=Cm(.7)
    s.different_first_page_header_footer=True;element(s._sectPr,'w:bidi')
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='B Nazanin';st.font.size=Pt(14);st.font.italic=False
        st.font.color.rgb=RGBColor(0,0,0);pr=st.element.get_or_add_rPr()
        for k in ['ascii','hAnsi','cs','eastAsia']:pr.rFonts.set(qn('w:'+k),'B Nazanin')
        element(pr,'w:szCs',{'w:val':'28'})
        for child in list(pr.findall(qn('w:iCs'))):pr.remove(child)
        element(pr,'w:iCs',{'w:val':'0'})
        ppr=st.element.get_or_add_pPr()
        for tag in ['w:pBdr','w:shd']:
            for child in list(ppr.findall(qn(tag))):ppr.remove(child)
    identity=university_mark(doc,'configs/current_report_identity.json')
    para(doc,identity['university_fa'],18,True,8)
    para(doc,identity['course_fa'],15,after=16)
    para(doc,'ارزیابی کنترل‌شده‌ی قطعه‌بندی',24,True,6,'Title')
    para(doc,'و تحلیل تشخیصی بازیابی شاهد',23,True,6,'Subtitle')
    para(doc,'در حافظه‌ی گفت‌وگویی',23,True,20,'Subtitle')
    para(doc,'پیش‌نویس یکپارچه‌ی پایان‌نامه',16,after=16)
    para(doc,'پژوهشگر  '+identity['author_fa'],16,True,10)
    para(doc,'استاد راهنمای دانشگاهی  '+identity['university_supervisor_fa'],15,after=10)
    para(doc,'راهنمای آموزشی علوم اعصاب شناختی  دکتر آرش فکری',14,after=7)
    para(doc,'فردیت FARDIAT مجتمع فوق تخصصی مغز و شناخت',14,after=7)
    p=para(doc,after=10);r=p.add_run();r.add_picture(str(ROOT/identity['educational_institute_logo']),width=Inches(.65))
    for d in r._r.xpath('.//wp:docPr'):d.set('descr',identity['educational_institute_fa'])
    para(doc,'۲۷ اکتبر ۲۰۲۶',14,after=12)
    para(doc,'github.com/yousofab/event-memory-retrieval',10.5,after=0)
    doc.core_properties.title='ارزیابی کنترل‌شده‌ی قطعه‌بندی و تحلیل تشخیصی بازیابی شاهد در حافظه‌ی گفت‌وگویی'
    doc.core_properties.author=identity['author_fa']
    doc.core_properties.subject='Integrated BSc thesis review draft; experimental scope through Phase19'
    footer=s.footer.paragraphs[0];rtl(footer);runs(footer,'پایان‌نامه‌ی کارشناسی  |  پیش‌نویس یکپارچه  |  ',10)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    markdown(doc,FOLDER/'front_matter_fa.md',anchor='abstractFA')
    markdown(doc,FOLDER/'abstract_en.md',anchor='abstractEN',english=True)
    items=contents(doc,page_map)
    for name in CHAPTERS:markdown(doc,FOLDER/'chapters'/name,anchor='ch'+name[:2])
    markdown(doc,FOLDER/'references_fa.md',anchor='externalSources')
    markdown(doc,FOLDER/'internal_sources_fa.md',anchor='internalSources')
    markdown(doc,FOLDER/'appendices_fa.md')
    out=ROOT/'deliverables/thesis_integrated_fa.docx';doc.save(out)
    print(out)
    return items


def main():
    p=argparse.ArgumentParser();p.add_argument('--page-map',type=Path);args=p.parse_args()
    build(json.loads(args.page_map.read_text()) if args.page_map else {})


if __name__=='__main__':main()
