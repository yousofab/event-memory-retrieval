"""Create the Persian RTL B Nazanin report from the audited Markdown source.

Report dependencies are separate from the standard-library experiment.
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm,Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import package_phase17 as helpers
from package_phase17 import element,rtl,runfont
from report_identity import university_mark

ROOT=Path(__file__).resolve().parents[1]
FONT='B Nazanin'
helpers.LATIN='DejaVu Sans'


def runs(p,text,size=14,bold=False,color='000000'):
    text=text.replace('**','').replace('`','')
    # The original B Nazanin does not reliably render U+066B. Keep decimal
    # digits and separator together in a font with all required glyphs.
    for part in re.split(r'([۰-۹]+٫[۰-۹]+|[\x21-\x7e]+(?:[ \t]+[\x21-\x7e]+)*)',text):
        if part:
            latin=bool(re.fullmatch(r'[\x20-\x7e]+|[۰-۹]+٫[۰-۹]+',part))
            runfont(p.add_run(part),size,bold,latin,color)


def para(doc,text='',size=14,bold=False,after=6,style=None):
    p=doc.add_paragraph(style=style);rtl(p)
    p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.10
    p.paragraph_format.widow_control=True
    runs(p,text,size,bold)
    return p


def heading(doc,text,page=False):
    p=para(doc,text,17,True,8,'Heading 1')
    p.paragraph_format.space_before=Pt(9)
    p.paragraph_format.page_break_before=page
    p.paragraph_format.keep_with_next=True


def table(doc,lines,number):
    names={1:'داده و دانه‌بندی مرجع',2:'نتیجه‌ی اصلی در سقف ۱۰۲۴ واژه',3:'دسترسی و امکان و پوشش کامل',4:'ممیزی صحت عددی'}
    p=para(doc,f'جدول {number} '+names[number],12,True,5)
    p.paragraph_format.keep_with_next=True
    records=[[c.strip() for c in line.strip('|').split('|')] for line in lines if not re.fullmatch(r'[|:\-\s]+',line)]
    n=len(records[0]);widths=[2.9,1.8,1.8] if n==3 else [2.3,1.35,1.35,1.5]
    t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.RIGHT;t.autofit=False
    element(t._tbl.tblPr,'w:bidiVisual');element(t._tbl.tblPr,'w:tblLayout',{'w:type':'fixed'})
    borders=element(t._tbl.tblPr,'w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        element(borders,'w:'+side,{'w:val':'single','w:sz':'4','w:color':'D9D9D9'})
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,row in enumerate(records):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=t.rows[i]._tr.get_or_add_trPr();element(trpr,'w:cantSplit')
        if i==0:element(trpr,'w:tblHeader')
        for j,text in enumerate(row):
            c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();element(pr,'w:shd',{'w:fill':'264B65' if i==0 else ('F1F5F7' if i%2 else 'FFFFFF')})
            margins=element(pr,'w:tcMar')
            for side in ['top','bottom','left','right']:element(margins,'w:'+side,{'w:w':'85','w:type':'dxa'})
            p=c.paragraphs[0];rtl(p)
            p.paragraph_format.keep_with_next=i<len(records)-1
            p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.08
            if j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            runs(p,text,12.5,i==0,'FFFFFF' if i==0 else '000000')
    para(doc,'',3,after=3)


def build():
    doc=Document();s=doc.sections[0]
    s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.top_margin,s.bottom_margin=Cm(1.8),Cm(1.8)
    s.left_margin=s.right_margin=Cm(2.2);s.footer_distance=Cm(.7)
    element(s._sectPr,'w:bidi')
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name=FONT;st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
        st.font.italic=False
        pr=st.element.get_or_add_rPr()
        for k in ['ascii','hAnsi','cs','eastAsia']:pr.rFonts.set(qn('w:'+k),FONT)
        element(pr,'w:szCs',{'w:val':'28'});element(pr,'w:iCs',{'w:val':'0'})
        ppr=st.element.get_or_add_pPr()
        for k in ['w:pBdr','w:shd']:
            for child in list(ppr.findall(qn(k))):ppr.remove(child)
    doc.core_properties.author='Yousef Abolfathi Roozbahani'
    doc.core_properties.title='دسترسی به جلسه و پوشش شاهد در حافظه‌ی گفت‌وگویی'
    doc.core_properties.subject='Phase19 session access and exact evidence audit'
    footer=s.footer.paragraphs[0];rtl(footer)
    runs(footer,'پژوهش حافظه‌ی رویدادی  |  مرحله‌ی ۱۹  |  ',10)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

    identity=university_mark(doc)
    para(doc,identity['university_fa'],18,True,8)
    para(doc,identity['course_fa'],15,after=14)
    para(doc,'گزارش پژوهشی مرحله‌ی ۱۹',25,True,14,'Title')
    para(doc,'دسترسی به جلسه و پوشش شاهد',24,True,4,'Subtitle')
    para(doc,'در حافظه‌ی گفت‌وگویی',22,True,18,'Subtitle')
    para(doc,'ممیزی کنترل‌شده با انتظار ترکیبی دقیق و ارزیابی روی منبع تازه',16,after=18)
    para(doc,'پژوهشگر  یوسف ابوالفتحی روزبهانی',16,True,10)
    para(doc,'استاد راهنمای دانشگاهی  استاد امیر نوری',15,after=8)
    para(doc,'راهنمای آموزشی علوم اعصاب شناختی  دکتر آرش فکری',15,after=8)
    para(doc,'فردیت  FARDIAT  مجتمع فوق تخصصی مغز و شناخت',15,after=12)
    if (ROOT/'assets/fardiat-logo.jpg').exists():
        p=para(doc,after=8);p.add_run().add_picture(str(ROOT/'assets/fardiat-logo.jpg'),width=Inches(.80))
    para(doc,'۲۵ اکتبر ۲۰۲۶',14,after=12)
    para(doc,'نتیجه‌ی این مرحله  افزایش دسترسی به جلسه، پوشش کافی شاهد را تضمین نکرد. داده، حدود ادعا و تحلیل تکمیلی اکتشافی در گزارش تفکیک شده‌اند.',13.5,after=12)
    para(doc,'github.com/yousofab/event-memory-retrieval',10.5,after=0)
    sections=re.split(r'^## (.+)$',(ROOT/'docs/phase19.md').read_text(),flags=re.M)
    table_number=0
    for title,body in zip(sections[1::2],sections[2::2]):
        page=title=='چکیده' or title.startswith(('۲ ','۳ ','۴ ','۸ ','۹ ','۱۰ ','منابع'))
        heading(doc,title,page)
        lines=body.strip().splitlines();i=0
        while i<len(lines):
            line=lines[i].strip()
            if not line or line.startswith('<'):
                i+=1;continue
            if line.startswith('|'):
                block=[]
                while i<len(lines) and lines[i].strip().startswith('|'):
                    block.append(lines[i].strip());i+=1
                table_number+=1;table(doc,block,table_number);continue
            image=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
            if image:
                p=para(doc,after=2);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1
                run=p.add_run();run.add_picture(str((ROOT/'docs'/image[2]).resolve()),width=Inches(6.1))
                for drawing in run._r.xpath('.//wp:docPr'):
                    drawing.set('descr',image[1])
                para(doc,image[1],11.5,after=8)
            elif re.match(r'^\[R\d+\]',line):
                p=doc.add_paragraph();rtl(p,False);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=1.1
                runfont(p.add_run(line),10,latin=True)
            else:para(doc,line)
            i+=1
    out=ROOT/'deliverables/phase19_research_report_fa.docx';out.parent.mkdir(parents=True,exist_ok=True)
    doc.save(out);print(out)


if __name__=='__main__':build()
